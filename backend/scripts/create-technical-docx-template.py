"""Creates the non-legal DOCX used to verify the parameterized generation pipeline."""
from pathlib import Path
import sys
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from docx.shared import RGBColor

output = Path(sys.argv[1])
document = Document()
section = document.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = section.bottom_margin = Inches(0.85)
section.left_margin = section.right_margin = Inches(0.9)

styles = document.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(11)
title = document.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_run = title.add_run("Technické ověření generování dokumentu")
title_run.bold = True
title_run.font.name = "Aptos Display"
title_run.font.size = Pt(20)
title_run.font.color.rgb = RGBColor(0, 0, 0)
title.paragraph_format.space_after = Pt(18)
intro = document.add_paragraph("Tento dokument slouží pouze k ověření bezpečné generační a SharePoint pipeline DeveloCRM.")
intro.paragraph_format.space_after = Pt(18)

table = document.add_table(rows=0, cols=2)
table.style = "Table Grid"
for label, value in (
    ("Projekt", "{{project.name}}"),
    ("Kód projektu", "{{project.code}}"),
    ("Jednotka", "{{unit.code}}"),
    ("Klient", "{{buyer.name}}"),
    ("Datum generování", "{{generation.date}}"),
    ("Aktuální cena", "{{unit.totalPrice}}"),
):
    cells = table.add_row().cells
    cells[0].text = label
    cells[1].text = value
    cells[0].paragraphs[0].runs[0].bold = True

document.add_paragraph("Nejde o smlouvu ani jiný právní dokument.")
output.parent.mkdir(parents=True, exist_ok=True)
document.save(output)
