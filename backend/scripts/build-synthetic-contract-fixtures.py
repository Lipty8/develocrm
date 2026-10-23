"""Create explicitly non-legal fixtures; never read or modify legal originals."""
import argparse
import hashlib
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

parser = argparse.ArgumentParser()
parser.add_argument("tokens", type=Path)
parser.add_argument("output", type=Path)
args = parser.parse_args()
schemas = json.loads(args.tokens.read_text())
args.output.mkdir(parents=True, exist_ok=True)
manifest = {"syntheticOnly": True, "legalContent": False, "files": {}}
for kind in ("rs", "sbk"):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.75)
    for name in ("Normal", "Title"):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(11 if name == "Normal" else 18)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(6)
    doc.add_paragraph(f"Technický test toku {kind.upper()}", "Title")
    doc.add_paragraph("Neplatný technický soubor pro ověření přenosu údajů DeveloCRM. "
                      "Neobsahuje právní smluvní text a nezakládá závazek.")
    doc.add_paragraph("TECH72_SYNTHETIC_ONLY")
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width, table.columns[1].width = Inches(2.55), Inches(4.45)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Údaj", "Hodnota"
    repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for token in schemas[kind]:
        cells = table.add_row().cells
        cells[0].text = token
        cells[1].text = "{{" + token + "}}"
    for index, row in enumerate(table.rows):
        for column, cell in enumerate(row.cells):
            cell.width = Inches(2.55 if column == 0 else 4.45)
            props = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right"):
                border = OxmlElement("w:" + edge)
                for key, value in {"val": "single", "sz": "4", "color": "D9D9D9"}.items():
                    border.set(qn("w:" + key), value)
                borders.append(border)
            props.append(borders)
            margins = OxmlElement("w:tcMar")
            for edge in ("top", "left", "bottom", "right"):
                margin = OxmlElement("w:" + edge)
                margin.set(qn("w:w"), "80")
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            props.append(margins)
            if index == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "E7EEF5")
                props.append(shading)
                for run in cell.paragraphs[0].runs:
                    run.bold = True
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
    doc.core_properties.author = "DeveloCRM technical test"
    doc.core_properties.title = f"Technický test toku {kind.upper()}"
    output = args.output / f"TECH72-{kind}-synthetic-template.docx"
    doc.save(output)
    manifest["files"][kind] = {"filename": output.name, "sha256": hashlib.sha256(output.read_bytes()).hexdigest(),
                                "tokens": schemas[kind]}
(args.output / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n")
print(json.dumps({"syntheticOnly": True, "files": [v["filename"] for v in manifest["files"].values()]}))
